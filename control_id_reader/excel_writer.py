from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from control_id_reader.utils import (
    MESES_ABREV,
    formatar_cpf,
    media_horas,
    mes_ano_predominante,
    nome_para_arquivo,
    parse_data_ponto,
    somar_duracoes,
)


def generate_suggested_filename(dados, todos_funcionarios, input_count, now=None):
    """Gera o nome sugerido para o arquivo Excel com base no cenario."""
    now = now or datetime.now()
    num_funcionarios = len(todos_funcionarios)

    if num_funcionarios == 1:
        info = list(todos_funcionarios.values())[0]
        nome_curto = _nome_curto(info.get("nome", "Funcionario"))
        dias_funcionario = info.get("_dias", dados)

        if input_count == 1:
            mes_ano = mes_ano_predominante(dias_funcionario)
            sufixo = f"_{mes_ano}" if mes_ano else ""
            return f"Ponto_{nome_curto}{sufixo}.xlsx"
        return f"Ponto_{nome_curto}.xlsx"

    mes_ano = MESES_ABREV[now.month] + str(now.year)
    return f"Ponto_{num_funcionarios}Funcionarios_{mes_ano}.xlsx"


def build_workbook(dados, todos_funcionarios, inconsistencias=None):
    """Cria e retorna um Workbook formatado com os dados de ponto."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ponto"

    funcionario_unico = len(todos_funcionarios) == 1

    fonte_cab = Font(bold=True, color="FFFFFF", size=10)
    fill_cab = PatternFill("solid", fgColor="1F4E79")
    fill_info = PatternFill("solid", fgColor="D6E4F0")
    fonte_info = Font(bold=True, size=10, color="1F4E79")
    fonte_label = Font(bold=True, size=9, color="2C3E50")
    alinhamento = Alignment(vertical="center", horizontal="left")
    alinhamento_centro = Alignment(vertical="center", horizontal="center")
    borda_fina = Border(
        bottom=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
    )

    linha_atual = 1

    def chave_ordenacao(dia):
        nome = dia.get("nome", "")
        return (nome, parse_data_ponto(dia.get("dia", "")))

    dados = sorted(dados, key=chave_ordenacao)

    if funcionario_unico:
        info = list(todos_funcionarios.values())[0]
        campos = [
            ("Empresa", info.get("empresa", "")),
            ("Nome", info.get("nome", "")),
            ("CPF", formatar_cpf(info.get("cpf", ""))),
            ("PIS/PASEP", info.get("pis", "")),
            ("Matrícula", info.get("matricula", "")),
            ("Admissão", info.get("admissao", "")),
            ("Cargo", info.get("cargo", "")),
            ("Departamento", info.get("departamento", "")),
            ("Centro de Custo", info.get("centro_custo", "")),
        ]
        campos = [(k, v) for k, v in campos if v]

        n_colunas_cab = 10

        ws.merge_cells(start_row=linha_atual, start_column=1, end_row=linha_atual, end_column=n_colunas_cab)
        cel = ws.cell(row=linha_atual, column=1, value="DADOS DO FUNCIONÁRIO")
        cel.font = fonte_cab
        cel.fill = fill_cab
        cel.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[linha_atual].height = 20
        linha_atual += 1

        for i in range(0, len(campos), 2):
            par = campos[i:i + 2]
            for col_offset, (label, valor) in enumerate(par):
                col_l = 1 + col_offset * 5
                col_v = col_l + 1
                col_fim = min(col_l + 4, n_colunas_cab)

                cel_l = ws.cell(row=linha_atual, column=col_l, value=label.upper() + ":")
                cel_l.font = fonte_label
                cel_l.fill = fill_info
                cel_l.alignment = alinhamento

                ws.merge_cells(start_row=linha_atual, start_column=col_v, end_row=linha_atual, end_column=col_fim)
                cel_v = ws.cell(row=linha_atual, column=col_v, value=valor)
                cel_v.font = fonte_info
                cel_v.fill = fill_info
                cel_v.alignment = alinhamento

            ultimo_col_usado = 1 + len(par) * 5 - 1
            for col_r in range(ultimo_col_usado + 1, n_colunas_cab + 1):
                ws.cell(row=linha_atual, column=col_r).fill = fill_info

            ws.row_dimensions[linha_atual].height = 18
            linha_atual += 1

        for col_r in range(1, n_colunas_cab + 1):
            ws.cell(row=linha_atual, column=col_r).fill = PatternFill("solid", fgColor="FFFFFF")
        ws.row_dimensions[linha_atual].height = 6
        linha_atual += 1

    if funcionario_unico:
        cabecalhos_visiveis = [
            "DIA", "MARCAÇÕES", "ENT. 1", "SAÍ. 1",
            "ENT. 2", "SAÍ. 2", "ENT. 3", "SAÍ. 3", "DURAÇÃO", "CH"
        ]
        valores_keys = [
            "dia", "marcacoes", "ent1", "sai1",
            "ent2", "sai2", "ent3", "sai3", "duracao", "ch"
        ]
        larguras = [18, 26, 10, 10, 10, 10, 10, 10, 10, 8]
    else:
        cabecalhos_visiveis = [
            "FUNCIONÁRIO", "CPF", "DIA", "MARCAÇÕES", "ENT. 1", "SAÍ. 1",
            "ENT. 2", "SAÍ. 2", "ENT. 3", "SAÍ. 3", "DURAÇÃO", "CH"
        ]
        valores_keys = [
            "nome", "cpf", "dia", "marcacoes", "ent1", "sai1",
            "ent2", "sai2", "ent3", "sai3", "duracao", "ch"
        ]
        larguras = [30, 16, 18, 26, 10, 10, 10, 10, 10, 10, 10, 8]

    fill_par = PatternFill("solid", fgColor="EBF5FB")
    fill_impar = PatternFill("solid", fgColor="FFFFFF")

    cab_resumo = [
        "NOME", "CPF", "PIS/PASEP", "CARGO", "ADMISSÃO",
        "MATRÍCULA", "DEPARTAMENTO", "CENTRO DE CUSTO",
        "DIAS TRABALHADOS", "DIAS FALTADOS", "DIAS DE TRABALHO TOTAIS",
        "MÉDIA HORAS/DIA"
    ]
    larg_resumo = [35, 16, 14, 25, 12, 12, 22, 18, 16, 14, 22, 16]

    def linha_resumo(info, dias_lista):
        trab = dias_trabalhados(dias_lista)
        total_trabalhados = len(trab)
        total_faltados = len(dias_faltados(dias_lista))
        total_trabalho = len(dias_trabalho_totais(dias_lista))
        total_horas = somar_duracoes([d.get("duracao", "") for d in trab])
        media = media_horas(total_horas, total_trabalhados)
        return [
            info.get("nome", ""),
            formatar_cpf(info.get("cpf", "")),
            info.get("pis", ""),
            info.get("cargo", ""),
            info.get("admissao", ""),
            info.get("matricula", ""),
            info.get("departamento", ""),
            info.get("centro_custo", ""),
            total_trabalhados,
            total_faltados,
            total_trabalho,
            media,
        ]

    fill_resumo_cab = PatternFill("solid", fgColor="2E86C1")
    n_colunas_resumo = len(cab_resumo)

    if funcionario_unico:
        ws.merge_cells(start_row=linha_atual, start_column=1, end_row=linha_atual, end_column=n_colunas_resumo)
        cel = ws.cell(row=linha_atual, column=1, value="RESUMO")
        cel.font = fonte_cab
        cel.fill = fill_resumo_cab
        cel.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[linha_atual].height = 20
        linha_atual += 1

        for col, cab in enumerate(cab_resumo, 1):
            cel = ws.cell(row=linha_atual, column=col, value=cab)
            cel.font = fonte_cab
            cel.fill = fill_resumo_cab
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda_fina
        ws.row_dimensions[linha_atual].height = 18
        linha_atual += 1

        info_u = list(todos_funcionarios.values())[0]
        vals_r = linha_resumo(info_u, info_u.get("_dias", dados))
        for col, valor in enumerate(vals_r, 1):
            cel = ws.cell(row=linha_atual, column=col, value=valor)
            cel.fill = fill_par
            cel.alignment = alinhamento
            cel.border = borda_fina
        linha_atual += 1

        for col_r in range(1, n_colunas_resumo + 1):
            ws.cell(row=linha_atual, column=col_r).fill = PatternFill("solid", fgColor="FFFFFF")
        ws.row_dimensions[linha_atual].height = 6
        linha_atual += 1
    else:
        ws_resumo = wb.create_sheet("Resumo")

        for col, cab in enumerate(cab_resumo, 1):
            cel = ws_resumo.cell(row=1, column=col, value=cab)
            cel.font = fonte_cab
            cel.fill = fill_cab
            cel.alignment = Alignment(horizontal="center", vertical="center")
        ws_resumo.row_dimensions[1].height = 20

        for row_r, (_cpf, info) in enumerate(todos_funcionarios.items(), 2):
            vals_r = linha_resumo(info, info.get("_dias", []))
            fill = fill_par if row_r % 2 == 0 else fill_impar
            for col, valor in enumerate(vals_r, 1):
                cel = ws_resumo.cell(row=row_r, column=col, value=valor)
                cel.fill = fill
                cel.alignment = alinhamento
                cel.border = borda_fina

        for col, larg in enumerate(larg_resumo, 1):
            ws_resumo.column_dimensions[get_column_letter(col)].width = larg

    for col, cab in enumerate(cabecalhos_visiveis, 1):
        cel = ws.cell(row=linha_atual, column=col, value=cab)
        cel.font = fonte_cab
        cel.fill = fill_cab
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border = borda_fina
    ws.row_dimensions[linha_atual].height = 20
    linha_atual += 1

    for i, dia in enumerate(dados):
        fill = fill_par if i % 2 == 0 else fill_impar
        col_dia = 1 if funcionario_unico else 3
        for col, key in enumerate(valores_keys, 1):
            cel = ws.cell(row=linha_atual, column=col, value=dia.get(key, ""))
            cel.fill = fill
            cel.alignment = alinhamento if col == col_dia else alinhamento_centro
            cel.border = borda_fina
        linha_atual += 1

    for col, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col)].width = largura

    if inconsistencias:
        _adicionar_aba_inconsistencias(wb, inconsistencias)

    resumo_ch = build_ch_summary_rows(todos_funcionarios)
    if resumo_ch:
        _adicionar_aba_resumo_ch(wb, resumo_ch)

    _ordenar_abas(wb)

    return wb


def save_excel_file(excel_path, dados, todos_funcionarios, inconsistencias=None):
    """Cria e salva o arquivo Excel."""
    wb = build_workbook(dados, todos_funcionarios, inconsistencias=inconsistencias)
    wb.save(excel_path)


def dias_trabalhados(dias_lista):
    """Filtra dias que possuem ao menos uma marcacao de entrada ou saida."""
    return [
        dia for dia in dias_lista
        if any(dia.get(k, "").strip() for k in ("ent1", "sai1", "ent2", "sai2", "ent3", "sai3"))
    ]


def dias_trabalho_totais(dias_lista):
    """Filtra dias de trabalho, definidos pela coluna CH preenchida."""
    return [dia for dia in dias_lista if dia.get("ch", "").strip()]


def dias_faltados(dias_lista):
    """Filtra dias de trabalho sem nenhuma marcacao de ponto."""
    return [
        dia for dia in dias_trabalho_totais(dias_lista)
        if not any(dia.get(k, "").strip() for k in ("ent1", "sai1", "ent2", "sai2", "ent3", "sai3"))
    ]


def build_ch_summary_rows(todos_funcionarios):
    """Monta linhas de resumo por funcionario e codigo CH."""
    rows = []
    for _chave, info in sorted(todos_funcionarios.items(), key=lambda item: item[1].get("nome", "")):
        dias_por_ch = {}
        for dia in info.get("_dias", []):
            ch = dia.get("ch", "").strip()
            if not ch:
                continue
            dias_por_ch.setdefault(ch, []).append(dia)

        horarios_por_codigo = _horarios_por_codigo(info.get("horarios_contratuais", []))

        for ch, dias in sorted(dias_por_ch.items()):
            rows.append({
                "funcionario": info.get("nome", ""),
                "cpf": formatar_cpf(info.get("cpf", "")),
                "ch": ch,
                "horario": horarios_por_codigo.get(ch, ""),
                "dias_trabalhados": len(dias_trabalhados(dias)),
                "dias_faltados": len(dias_faltados(dias)),
                "dias_trabalho_totais": len(dias_trabalho_totais(dias)),
            })

    return rows


def _nome_curto(nome_completo):
    partes = nome_para_arquivo(nome_completo or "Funcionario").split("_")
    return f"{partes[0]}_{partes[-1]}" if len(partes) > 1 else partes[0]


def _horarios_por_codigo(horarios_contratuais):
    return {
        str(horario.get("codigo", "")).strip(): _formatar_horario_contratual(horario)
        for horario in horarios_contratuais
        if str(horario.get("codigo", "")).strip()
    }


def _formatar_horario_contratual(horario):
    pares = []
    for idx in range(1, 7):
        ent = str(horario.get(f"ent{idx}", "") or "").strip()
        sai = str(horario.get(f"sai{idx}", "") or "").strip()
        if ent or sai:
            pares.append(f"{ent}-{sai}".strip("-"))
    return " / ".join(pares)


def _adicionar_aba_resumo_ch(wb, rows):
    ws = wb.create_sheet("Resumo CH")

    headers = [
        "FUNCIONÁRIO", "CPF", "CH", "HORÁRIO CONTRATUAL",
        "DIAS TRABALHADOS", "DIAS FALTADOS", "DIAS DE TRABALHO TOTAIS"
    ]
    larguras = [34, 16, 10, 36, 18, 14, 22]

    fonte_cab = Font(bold=True, color="FFFFFF", size=10)
    fill_cab = PatternFill("solid", fgColor="1F4E79")
    fill_par = PatternFill("solid", fgColor="EBF5FB")
    fill_impar = PatternFill("solid", fgColor="FFFFFF")
    alinhamento = Alignment(vertical="center", horizontal="left", wrap_text=True)
    alinhamento_centro = Alignment(vertical="center", horizontal="center")
    borda_fina = Border(bottom=Side(style="thin", color="BFBFBF"))

    for col, header in enumerate(headers, 1):
        cel = ws.cell(row=1, column=col, value=header)
        cel.font = fonte_cab
        cel.fill = fill_cab
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border = borda_fina

    for row_idx, row in enumerate(rows, 2):
        fill = fill_par if row_idx % 2 == 0 else fill_impar
        valores = [
            row["funcionario"],
            row["cpf"],
            row["ch"],
            row["horario"],
            row["dias_trabalhados"],
            row["dias_faltados"],
            row["dias_trabalho_totais"],
        ]
        for col, valor in enumerate(valores, 1):
            cel = ws.cell(row=row_idx, column=col, value=valor)
            cel.fill = fill
            cel.alignment = alinhamento if col in (1, 4) else alinhamento_centro
            cel.border = borda_fina

    for col, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col)].width = largura


def _adicionar_aba_inconsistencias(wb, inconsistencias):
    ws = wb.create_sheet("Inconsistências")

    headers = ["SEVERIDADE", "FUNCIONÁRIO", "CPF", "DIA", "CAMPO", "MENSAGEM"]
    larguras = [14, 34, 16, 18, 18, 64]

    fonte_cab = Font(bold=True, color="FFFFFF", size=10)
    fill_cab = PatternFill("solid", fgColor="C0392B")
    fill_alta = PatternFill("solid", fgColor="FADBD8")
    fill_atencao = PatternFill("solid", fgColor="FCF3CF")
    fill_info = PatternFill("solid", fgColor="EBF5FB")
    alinhamento = Alignment(vertical="top", horizontal="left", wrap_text=True)
    borda_fina = Border(bottom=Side(style="thin", color="BFBFBF"))

    for col, header in enumerate(headers, 1):
        cel = ws.cell(row=1, column=col, value=header)
        cel.font = fonte_cab
        cel.fill = fill_cab
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border = borda_fina

    for row, issue in enumerate(inconsistencias, 2):
        fill = _fill_para_severidade(issue.get("severidade", ""), fill_alta, fill_atencao, fill_info)
        valores = [
            issue.get("severidade", ""),
            issue.get("funcionario", ""),
            issue.get("cpf", ""),
            issue.get("dia", ""),
            issue.get("campo", ""),
            issue.get("mensagem", ""),
        ]
        for col, valor in enumerate(valores, 1):
            cel = ws.cell(row=row, column=col, value=valor)
            cel.fill = fill
            cel.alignment = alinhamento
            cel.border = borda_fina

    for col, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col)].width = largura


def _ordenar_abas(wb):
    ordem_preferida = ["Ponto", "Inconsistências", "Resumo CH", "Resumo"]
    abas_por_nome = {ws.title: ws for ws in wb.worksheets}
    abas_ordenadas = [abas_por_nome[nome] for nome in ordem_preferida if nome in abas_por_nome]
    abas_ordenadas.extend(ws for ws in wb.worksheets if ws.title not in ordem_preferida)
    wb._sheets = abas_ordenadas


def _fill_para_severidade(severidade, fill_alta, fill_atencao, fill_info):
    if severidade == "Alta":
        return fill_alta
    if severidade == "Atenção":
        return fill_atencao
    return fill_info
