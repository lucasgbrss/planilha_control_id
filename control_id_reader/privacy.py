import base64
import json
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from cryptography.fernet import Fernet, InvalidToken
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from openpyxl import load_workbook


KEY_FORMAT = "ControlIDReaderPrivacyKey"
KEY_VERSION = 1
KDF_ITERATIONS = 390_000

SENSITIVE_HEADERS = {
    "FUNCIONARIO": "funcionario",
    "NOME": "funcionario",
    "CPF": "cpf",
    "PIS": "pis",
    "PISPASEP": "pis",
    "MATRICULA": "matricula",
    "EMPRESA": "empresa",
    "CNPJ": "cnpj",
}

KIND_PREFIX = {
    "funcionario": "FUNCIONARIO",
    "cpf": "CPF",
    "pis": "PIS",
    "matricula": "MATRICULA",
    "empresa": "EMPRESA",
    "cnpj": "CNPJ",
}


class PrivacyError(Exception):
    """Erro controlado nas rotinas de privacidade."""


def suggest_anonymized_path(excel_path):
    path = Path(excel_path)
    return path.with_name(f"{path.stem}_anonimizado{path.suffix}")


def suggest_restored_path(excel_path):
    path = Path(excel_path)
    stem = path.stem
    if stem.endswith("_anonimizado"):
        stem = stem[:-12]
    return path.with_name(f"{stem}_restaurado{path.suffix}")


def suggest_key_path(excel_path):
    path = Path(excel_path)
    return path.with_name(f"{path.stem}.cidkey")


def anonymize_excel_file(input_path, output_path, key_path, password):
    """Anonimiza dados sensiveis de uma planilha e salva uma chave criptografada."""
    if not password:
        raise PrivacyError("A senha da chave não pode ficar vazia.")

    wb = load_workbook(input_path)
    replacements, value_map = _build_anonymization_map(wb)

    if not replacements:
        raise PrivacyError("Nenhum dado sensível foi encontrado para anonimizar.")

    _apply_anonymization(wb, replacements, value_map)
    wb.save(output_path)

    payload = {
        "format": KEY_FORMAT,
        "version": KEY_VERSION,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "source_file": Path(input_path).name,
        "anonymized_file": Path(output_path).name,
        "replacements": replacements,
    }
    _write_encrypted_key(key_path, payload, password)
    return {
        "replacements": len(replacements),
        "unique_values": len(value_map),
        "output_path": str(output_path),
        "key_path": str(key_path),
    }


def restore_excel_file(input_path, output_path, key_path, password):
    """Restaura uma planilha anonimizada usando o arquivo de chave criptografado."""
    if not password:
        raise PrivacyError("A senha da chave não pode ficar vazia.")

    payload = _read_encrypted_key(key_path, password)
    replacements = payload.get("replacements", [])
    if not replacements:
        raise PrivacyError("A chave não contém dados para restaurar.")

    wb = load_workbook(input_path)
    restored = _apply_restore(wb, replacements)
    wb.save(output_path)

    return {
        "restored": restored,
        "output_path": str(output_path),
        "source_file": payload.get("source_file", ""),
    }


def _build_anonymization_map(wb):
    replacements = []
    by_kind_original = {}
    counters = {}

    def pseudonym_for(kind, original):
        key = (kind, original)
        if key not in by_kind_original:
            counters[kind] = counters.get(kind, 0) + 1
            by_kind_original[key] = f"{KIND_PREFIX[kind]}_{counters[kind]:03d}"
        return by_kind_original[key]

    seen_cells = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            header_columns = {}
            for cell in row:
                kind = _kind_for_header(cell.value)
                if kind:
                    header_columns[cell.column] = kind

                label_kind = _kind_for_label(cell.value)
                if label_kind:
                    target = ws.cell(row=cell.row, column=cell.column + 1)
                    _add_replacement(
                        replacements,
                        seen_cells,
                        ws.title,
                        target,
                        label_kind,
                        pseudonym_for,
                    )

            if header_columns:
                for data_row in ws.iter_rows(min_row=row[0].row + 1, max_row=ws.max_row):
                    if _row_is_empty(data_row):
                        break
                    for column, kind in header_columns.items():
                        cell = data_row[column - 1]
                        _add_replacement(
                            replacements,
                            seen_cells,
                            ws.title,
                            cell,
                            kind,
                            pseudonym_for,
                        )

    value_map = {item["original"]: item["anonymized"] for item in replacements}
    return replacements, value_map


def _add_replacement(replacements, seen_cells, sheet_name, cell, kind, pseudonym_for):
    value = _cell_text(cell.value)
    if not value or _is_header_like(value):
        return

    key = (sheet_name, cell.coordinate)
    if key in seen_cells:
        return

    anonymized = pseudonym_for(kind, value)
    replacements.append({
        "sheet": sheet_name,
        "cell": cell.coordinate,
        "kind": kind,
        "original": value,
        "anonymized": anonymized,
    })
    seen_cells.add(key)


def _apply_anonymization(wb, replacements, value_map):
    for item in replacements:
        ws = wb[item["sheet"]]
        ws[item["cell"]].value = item["anonymized"]

    text_values = _ordered_replacements(value_map.items())
    if not text_values:
        return

    for ws in wb.worksheets:
        sensitive_columns = _sensitive_columns(ws)
        for row in ws.iter_rows():
            for cell in row:
                if cell.column in sensitive_columns:
                    continue
                if isinstance(cell.value, str):
                    cell.value = _replace_text_tokens(cell.value, text_values)


def _apply_restore(wb, replacements):
    reverse_map = {item["anonymized"]: item["original"] for item in replacements}
    text_values = _ordered_replacements(reverse_map.items())
    restored = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                original_value = cell.value
                cell.value = _replace_text_tokens(cell.value, text_values)
                if cell.value != original_value:
                    restored += 1

    for item in replacements:
        if item["sheet"] in wb.sheetnames:
            ws = wb[item["sheet"]]
            if ws[item["cell"]].value == item["anonymized"]:
                ws[item["cell"]].value = item["original"]
                restored += 1

    return restored


def _sensitive_columns(ws):
    columns = set()
    for row in ws.iter_rows():
        for cell in row:
            if _kind_for_header(cell.value):
                columns.add(cell.column)
    return columns


def _replace_text_tokens(text, replacements):
    novo = text
    for original, anonymized in replacements:
        if original in novo:
            novo = novo.replace(original, anonymized)
    return novo


def _ordered_replacements(items):
    return sorted(
        (
            (str(original), str(anonymized))
            for original, anonymized in items
            if _safe_for_text_replacement(original) and anonymized
        ),
        key=lambda pair: len(pair[0]),
        reverse=True,
    )


def _safe_for_text_replacement(value):
    text = str(value or "").strip()
    return len(text) >= 5 or " " in text or any(char in text for char in ".-/")


def _kind_for_header(value):
    return SENSITIVE_HEADERS.get(_normalize_label(value))


def _kind_for_label(value):
    text = _cell_text(value)
    if not text.endswith(":"):
        return None
    return SENSITIVE_HEADERS.get(_normalize_label(text[:-1]))


def _normalize_label(value):
    text = _cell_text(value)
    text = unicodedata.normalize("NFKD", text)
    text = text.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9]", "", text).upper()


def _cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _row_is_empty(row):
    return not any(_cell_text(cell.value) for cell in row)


def _is_header_like(value):
    return bool(_kind_for_header(value) or _kind_for_label(value))


def _write_encrypted_key(path, payload, password):
    salt = os.urandom(16)
    token = _fernet(password, salt).encrypt(
        json.dumps(payload, ensure_ascii=False).encode("utf-8")
    )
    wrapper = {
        "format": KEY_FORMAT,
        "version": KEY_VERSION,
        "kdf": "PBKDF2HMAC-SHA256",
        "iterations": KDF_ITERATIONS,
        "salt": base64.urlsafe_b64encode(salt).decode("ascii"),
        "token": token.decode("ascii"),
    }
    Path(path).write_text(json.dumps(wrapper, indent=2, ensure_ascii=False), encoding="utf-8")


def _read_encrypted_key(path, password):
    try:
        wrapper = json.loads(Path(path).read_text(encoding="utf-8"))
        if wrapper.get("format") != KEY_FORMAT:
            raise PrivacyError("Arquivo de chave inválido.")
        salt = base64.urlsafe_b64decode(wrapper["salt"].encode("ascii"))
        token = wrapper["token"].encode("ascii")
        data = _fernet(password, salt).decrypt(token)
        payload = json.loads(data.decode("utf-8"))
    except InvalidToken as exc:
        raise PrivacyError("Senha incorreta ou arquivo de chave inválido.") from exc
    except (OSError, KeyError, ValueError, json.JSONDecodeError) as exc:
        raise PrivacyError("Não foi possível ler o arquivo de chave.") from exc

    if payload.get("format") != KEY_FORMAT:
        raise PrivacyError("Arquivo de chave inválido.")
    return payload


def _fernet(password, salt):
    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(),
        length=32,
        salt=salt,
        iterations=KDF_ITERATIONS,
    )
    key = base64.urlsafe_b64encode(kdf.derive(password.encode("utf-8")))
    return Fernet(key)
