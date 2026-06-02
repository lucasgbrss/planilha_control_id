import os
import zipfile
import email as email_lib
import pdfplumber
import pypdf
from bs4 import BeautifulSoup
from pathlib import Path
from openpyxl import Workbook
import re
import tkinter as tk
from tkinter import filedialog, messagebox
import customtkinter as ctk
from datetime import datetime
import threading
import json
import os


class PdfToExcelApp:
    """Aplicativo moderno para converter folhas de ponto em Excel."""

    def __init__(self, root):
        self.root = root
        self.root.title("Control ID Reader - Conversor de Ponto")
        self.root.geometry("820x680")
        self.root.resizable(True, True)
        self.root.minsize(700, 580)

        # Tema padrão
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        # Aplicar ícone da janela
        self._aplicar_icone()

        # Centralizar janela na tela
        self.centralizar_janela()

        # Variáveis de estado
        self.pdf_paths = []
        self.processamento_ativo = False
        self._tema_atual = "dark"

        # Diretórios padrão (serão sobrescritos pelas configurações salvas)
        self.dir_abrir_pdf = str(Path.home() / "Downloads")
        self.dir_salvar_excel = str(Path.home() / "Documents")

        # Criar interface
        self.criar_interface()

        # Carregar configurações salvas
        self.carregar_configuracoes()

    def _aplicar_icone(self):
        """Aplica o ícone da janela a partir do arquivo .ico na pasta do script."""
        try:
            ico_path = Path(__file__).parent / "control_id_reader.ico"
            if ico_path.exists():
                self.root.iconbitmap(str(ico_path))
        except Exception:
            pass  # Silencioso — ícone é cosmético, não deve travar o app

    def centralizar_janela(self):
        """Centraliza a janela na tela."""
        self.root.update_idletasks()
        largura = self.root.winfo_width()
        altura = self.root.winfo_height()
        tela_largura = self.root.winfo_screenwidth()
        tela_altura = self.root.winfo_screenheight()
        x = (tela_largura // 2) - (largura // 2)
        y = (tela_altura // 2) - (altura // 2)
        self.root.geometry(f"{largura}x{altura}+{x}+{y}")

    def alternar_tema(self):
        """Alterna entre dark e light mode."""
        if self._tema_atual == "dark":
            ctk.set_appearance_mode("light")
            self._tema_atual = "light"
            self.btn_tema.configure(text="☀️  Light")
        else:
            ctk.set_appearance_mode("dark")
            self._tema_atual = "dark"
            self.btn_tema.configure(text="🌙  Dark")
        self.root.after(50, self._atualizar_cores_log)

    def _atualizar_cores_log(self):
        """Sincroniza as cores do tk.Text e tk.Listbox com o tema CTk atual."""
        is_dark = self._tema_atual == "dark"
        bg      = "#1e1e1e" if is_dark else "#f0f0f0"
        fg      = "#d4d4d4" if is_dark else "#222222"
        sel_bg  = "#264f78" if is_dark else "#1f6aa5"

        self.log_text.configure(bg=bg, fg=fg,
                                insertbackground=fg, selectbackground=sel_bg)
        self.log_text.tag_config("info",       foreground="#d4d4d4" if is_dark else "#333333")
        self.log_text.tag_config("success",    foreground="#4ec94e" if is_dark else "#1e7e34")
        self.log_text.tag_config("error",      foreground="#f47070" if is_dark else "#c0392b")
        self.log_text.tag_config("processing", foreground="#f0c060" if is_dark else "#d68910")

        self.lista_arquivos.configure(
            bg=bg, fg=fg,
            selectbackground=sel_bg, selectforeground="white"
        )

    def criar_interface(self):
        """Cria todos os componentes da interface."""
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)

        # Container principal
        main_frame = ctk.CTkFrame(self.root, fg_color="transparent")
        main_frame.grid(row=0, column=0, sticky="nsew", padx=20, pady=20)
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(2, weight=1)  # log expande verticalmente

        # ── Cabeçalho ───────────────────────────────────────────
        header = ctk.CTkFrame(main_frame, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        header.columnconfigure(0, weight=1)

        ctk.CTkLabel(header, text="Control ID Reader",
                     font=ctk.CTkFont(family="Segoe UI", size=22, weight="bold")
                     ).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(header, text="Conversor de Folhas de Ponto para Excel",
                     font=ctk.CTkFont(family="Segoe UI", size=12),
                     text_color=("gray40", "gray70")
                     ).grid(row=1, column=0, sticky="w")

        # Botão tema no canto superior direito
        self.btn_tema = ctk.CTkButton(
            header, text="🌙  Dark", width=100, height=30,
            fg_color=("gray80", "gray25"), text_color=("gray10", "gray90"),
            hover_color=("gray70", "gray35"),
            command=self.alternar_tema
        )
        self.btn_tema.grid(row=0, column=1, rowspan=2, sticky="e")

        # Seções
        self.criar_secao_selecao(main_frame, 1)
        self.criar_secao_status(main_frame, 2)
        self.criar_secao_acoes(main_frame, 3)
    def criar_secao_selecao(self, parent, row):
        """Cria a seção para seleção de arquivos."""
        frame = ctk.CTkFrame(parent)
        frame.grid(row=row, column=0, sticky="ew", pady=(0, 8))
        frame.columnconfigure(0, weight=1)

        ctk.CTkLabel(frame, text="Seleção de Arquivos",
                     font=ctk.CTkFont(size=13, weight="bold")
                     ).grid(row=0, column=0, sticky="w", padx=14, pady=(10, 6))

        # Lista de arquivos (tk.Listbox ainda não tem substituto CTk nativo)
        list_frame = ctk.CTkFrame(frame, fg_color="transparent")
        list_frame.grid(row=1, column=0, sticky="ew", padx=12, pady=(0, 4))
        list_frame.columnconfigure(0, weight=1)

        self.lista_arquivos = tk.Listbox(
            list_frame, height=5, selectmode=tk.EXTENDED,
            font=("Segoe UI", 9), relief="flat", borderwidth=0,
            activestyle="none", selectbackground="#1f6aa5",
            selectforeground="white"
        )
        self.lista_arquivos.grid(row=0, column=0, sticky="ew")

        scrollbar = ctk.CTkScrollbar(list_frame, command=self.lista_arquivos.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.lista_arquivos.configure(yscrollcommand=scrollbar.set)

        # Botões
        btn_frame = ctk.CTkFrame(frame, fg_color="transparent")
        btn_frame.grid(row=2, column=0, sticky="w", padx=12, pady=(4, 4))

        ctk.CTkButton(btn_frame, text="+ Adicionar Arquivos", width=160,
                      fg_color="#27ae60", hover_color="#1e8449",
                      command=self.adicionar_pdfs
                      ).grid(row=0, column=0, padx=(0, 6))

        ctk.CTkButton(btn_frame, text="✕ Remover Selecionados", width=170,
                      fg_color="#e67e22", hover_color="#ca6f1e",
                      command=self.remover_selecionados
                      ).grid(row=0, column=1, padx=(0, 6))

        ctk.CTkButton(btn_frame, text="🗑  Limpar Tudo", width=130,
                      fg_color="#c0392b", hover_color="#a93226",
                      command=self.limpar_lista
                      ).grid(row=0, column=2)

        self.info_arquivos = ctk.CTkLabel(
            frame, text="Nenhum arquivo selecionado",
            font=ctk.CTkFont(size=11), text_color=("gray40", "gray60")
        )
        self.info_arquivos.grid(row=3, column=0, sticky="w", padx=14, pady=(0, 10))

    def criar_secao_status(self, parent, row):
        """Cria a seção de log/status."""
        frame = ctk.CTkFrame(parent)
        frame.grid(row=row, column=0, sticky="nsew", pady=(0, 8))
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(1, weight=1)

        ctk.CTkLabel(frame, text="Log de Processamento",
                     font=ctk.CTkFont(size=13, weight="bold")
                     ).grid(row=0, column=0, sticky="w", padx=14, pady=(10, 6))

        self.log_text = tk.Text(
            frame, height=9, wrap=tk.WORD,
            font=("Consolas", 9), relief="flat", borderwidth=0,
            padx=8, pady=6
        )
        self.log_text.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 12))

        log_scroll = ctk.CTkScrollbar(frame, command=self.log_text.yview)
        log_scroll.grid(row=1, column=1, sticky="ns", pady=(0, 12), padx=(0, 8))
        self.log_text.configure(yscrollcommand=log_scroll.set)

        self._atualizar_cores_log()

    def criar_secao_acoes(self, parent, row):
        """Cria a seção de botões de ação."""
        frame = ctk.CTkFrame(parent, fg_color="transparent")
        frame.grid(row=row, column=0, sticky="ew", pady=(4, 0))
        frame.columnconfigure(0, weight=1)

        # Barra de progresso
        self.progress_var = tk.DoubleVar()
        self.progress_bar = ctk.CTkProgressBar(frame, variable=self.progress_var,
                                               height=10)
        self.progress_bar.set(0)
        self.progress_bar.grid(row=0, column=0, columnspan=3, sticky="ew",
                               pady=(0, 10))

        # Botão principal
        self.btn_processar = ctk.CTkButton(
            frame, text="⚡  Gerar Planilha Excel",
            font=ctk.CTkFont(size=13, weight="bold"),
            height=40, fg_color="#2980b9", hover_color="#2471a3",
            command=self.iniciar_processamento
        )
        self.btn_processar.grid(row=1, column=0, sticky="w")

        # Botão secundário
        ctk.CTkButton(
            frame, text="💾  Salvar Configurações",
            height=40, width=180,
            fg_color=("gray75", "gray30"), text_color=("gray10", "gray90"),
            hover_color=("gray65", "gray40"),
            command=self.salvar_configuracoes
        ).grid(row=1, column=1, padx=(10, 0), sticky="w")

        ctk.CTkButton(
            frame, text="✂️  Separar PDF por Funcionário",
            height=40, width=220,
            fg_color="#7d3c98", hover_color="#6c3483",
            command=self.separar_pdf_por_funcionario
        ).grid(row=1, column=2, padx=(10, 0), sticky="w")

        self.status_final = ctk.CTkLabel(
            frame, text="",
            font=ctk.CTkFont(size=11), text_color=("gray40", "gray60")
        )
        self.status_final.grid(row=2, column=0, columnspan=3,
                               sticky="w", pady=(8, 0))

    def adicionar_pdfs(self):
        """Abre dialog para selecionar arquivos PDF ou MHTML."""
        arquivo_types = [
            ("Arquivos suportados", "*.pdf *.mhtml"),
            ("Arquivos PDF", "*.pdf"),
            ("Arquivos MHTML", "*.mhtml"),
        ]
        files = filedialog.askopenfilenames(
            title="Selecionar Arquivos PDF ou MHTML",
            filetypes=arquivo_types,
            initialdir=self.dir_abrir_pdf
        )

        if files:
            for f in files:
                if f not in self.pdf_paths:
                    self.pdf_paths.append(f)
                    self.lista_arquivos.insert(tk.END, Path(f).name)
            self.info_arquivos.configure(text=f"{len(self.pdf_paths)} arquivo(s) selecionado(s)")

            # Atualizar o último diretório usado
            self.dir_abrir_pdf = str(Path(files[0]).parent)
            self.salvar_configuracoes(silencioso=True)

    def remover_selecionados(self):
        """Remove arquivos selecionados da lista."""
        indices = self.lista_arquivos.curselection()
        if not indices:
            return

        for i in reversed(indices):
            self.lista_arquivos.delete(i)
            self.pdf_paths.pop(i)

        self.info_arquivos.configure(text=f"{len(self.pdf_paths)} arquivo(s) selecionado(s)")

    def limpar_lista(self):
        """Limpa todos os arquivos da lista."""
        self.pdf_paths.clear()
        self.lista_arquivos.delete(0, tk.END)
        self.info_arquivos.configure(text="Nenhum arquivo selecionado")

    def log(self, mensagem, tipo='info'):
        """Adiciona mensagem ao log de forma thread-safe via root.after."""
        timestamp = datetime.now().strftime("%H:%M:%S")
        def _inserir():
            self.log_text.insert(tk.END, f"[{timestamp}] {mensagem}\n", tipo)
            self.log_text.see(tk.END)
        self.root.after(0, _inserir)

    def _finalizar_processamento(self, todos_dias, total_pdfs, total_dias, todos_funcionarios):
        """Chamado na thread principal ao fim do processamento para abrir o filedialog."""
        self.processamento_ativo = False
        self.btn_processar.configure(state="normal")
        if todos_dias:
            self.salvar_excel(todos_dias, todos_funcionarios)
            self.log(f"\n✅ Processamento concluído!", "success")
            self.log(f"Total: {total_pdfs} arquivo(s), {total_dias} dia(s) registrado(s)", "success")
        else:
            self.log("\n⚠️ Nenhum dado foi extraído dos arquivos.", "error")

    def extrair_info_funcionario(self, texto):
        """Extrai informações do funcionário a partir do texto de um arquivo.
        Retorna (dict, motivo_falha). motivo_falha é None em caso de sucesso."""
        info = {}

        if not texto or not texto.strip():
            return info, "texto do arquivo está vazio"

        # EMPRESA e CNPJ na mesma linha
        match = re.search(r'EMPRESA:\s*(.+?)\s*CNPJ:\s*(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})',
                         texto, re.IGNORECASE)
        if match:
            info["empresa"] = match.group(1).strip()
            info["cnpj"] = match.group(2).strip()
        else:
            match = re.search(r'EMPRESA:\s*([^\n]+)', texto, re.IGNORECASE)
            if match:
                info["empresa"] = match.group(1).strip()
            match = re.search(r'CNPJ:\s*(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})', texto, re.IGNORECASE)
            if match:
                info["cnpj"] = match.group(1).strip()

        # ENDEREÇO
        match = re.search(r'ENDERE[Ç]O:\s*(.+?)(?:\n|$)', texto, re.IGNORECASE)
        if match:
            info["endereco"] = match.group(1).strip()

        # NOME
        match = re.search(r'NOME:\s*(.+?)\s+PIS/PASEP:', texto, re.IGNORECASE)
        if match:
            info["nome"] = match.group(1).strip()

        # PIS/PASEP
        match = re.search(r'PIS/PASEP:\s*(\d+)', texto, re.IGNORECASE)
        if match:
            info["pis"] = match.group(1).strip()

        # ADMISSÃO
        match = re.search(r'ADMISS[Ã]O:\s*(\d{2}/\d{2}/\d{4})', texto, re.IGNORECASE)
        if match:
            info["admissao"] = match.group(1).strip()

        # CPF
        match = re.search(r'CPF:\s*(\d+)', texto, re.IGNORECASE)
        if match:
            info["cpf"] = match.group(1).strip()

        # MATRÍCULA
        match = re.search(r'MATR[Í]CULA:\s*(\d+)', texto, re.IGNORECASE)
        if match:
            info["matricula"] = match.group(1).strip()

        # CENTRO DE CUSTO
        match = re.search(r'CENTRO DE CUSTO:\s*(\S+)', texto, re.IGNORECASE)
        if match:
            info["centro_custo"] = match.group(1).strip()

        # DEPARTAMENTO
        match = re.search(r'DEPARTAMENTO:\s*(\S+)', texto, re.IGNORECASE)
        if match:
            info["departamento"] = match.group(1).strip()

        # CARGO
        match = re.search(r'CARGO:\s*(.+?)(?:\n|$)', texto, re.IGNORECASE)
        if match:
            info["cargo"] = match.group(1).strip()

        # Fix 6: retornar motivo específico se nome não foi encontrado
        if not info.get("nome"):
            tem_empresa = "empresa" in info
            tem_cpf     = "cpf" in info
            if not tem_empresa and not tem_cpf:
                motivo = "nenhum campo reconhecido — formato do arquivo pode ser diferente do esperado"
            elif tem_empresa and not tem_cpf:
                motivo = "empresa identificada mas campo NOME não encontrado no padrão 'NOME: ... PIS/PASEP:'"
            else:
                motivo = "campo NOME não encontrado no padrão esperado"
            return info, motivo

        return info, None

    def extrair_dados_ponto(self, tabelas):
        """Extrai dados de ponto e horários contratuais de uma lista de tabelas.
        Retorna (dias, horarios_contratuais).

        Suporta múltiplos formatos de PDF — a detecção é semântica:
        - ent2/sai2 podem estar mesclados em [4] ou divididos por pdfplumber
        - CH (5 dígitos) e duração (HH:MM) são localizados por valor, não por índice fixo
        """
        dias = []
        horarios_contratuais = []
        capturando_horarios  = False

        TIME_RE = re.compile(r'^\d{1,2}:\d{2}$')
        CH_RE   = re.compile(r'^\d{5}$')
        # Padrão de linha de dia: deve começar com DD/MM/AA - DIA_SEMANA
        DAY_RE  = re.compile(r'^\d{2}/\d{2}/\d{2,4}\s*-\s*(SEG|TER|QUA|QUI|SEX|SAB|DOM)', re.IGNORECASE)

        MARCADORES_FIM = [
            "horários contratuais", "horarios contratuais",
            "código do horário",    "codigo do horario",
        ]

        def is_time(s): return bool(s and TIME_RE.match(str(s).strip()))
        def is_ch(s):   return bool(s and CH_RE.match(str(s).strip()))
        def cel(i):     return str(row[i]).strip() if i < len(row) and row[i] else ""

        for row in tabelas:
            if not row:
                continue

            linha_texto = " ".join(str(c) for c in row if c).lower()

            # ── Horários Contratuais ──────────────────────────────────────────
            if any(m in linha_texto for m in MARCADORES_FIM):
                capturando_horarios = True
                continue

            if capturando_horarios:
                codigo = cel(0)
                if codigo:
                    horario = {"codigo": codigo}
                    par = 1
                    col_h = 2
                    while col_h < len(row):
                        ent = cel(col_h)
                        sai = cel(col_h + 2) if col_h + 2 < len(row) else ""
                        if ent or sai:
                            horario[f"ent{par}"] = ent
                            horario[f"sai{par}"] = sai
                            par += 1
                        col_h += 4
                    horarios_contratuais.append(horario)
                continue

            # ── Linha de dia: deve ter formato DD/MM/AA - DIA ─────────────────
            if not (row[0] and DAY_RE.match(str(row[0]).strip())):
                continue

            dia = {
                "dia":       str(row[0]).strip(),
                "marcacoes": cel(1),
                "ent1":      cel(2),
                "sai1":      cel(3),
                "ent2": "", "sai2": "",
                "ent3": "", "sai3": "",
                "duracao":   "",
                "ch":        "",
            }

            # ── ent2 / sai2: detectar mescla e tempo dividido ─────────────────
            cel4_raw = cel(4)
            cel5_raw = cel(5)

            # Reconstituir tempo dividido pelo pdfplumber: "19:00 00" + ":05" → "19:00" + "00:05"
            if cel5_raw.startswith(":") and cel4_raw:
                parts4 = cel4_raw.split()
                if parts4:
                    hours_part  = parts4[-1]
                    cel4_raw    = " ".join(parts4[:-1])
                    cel5_raw    = f"{hours_part}{cel5_raw}"  # "00" + ":05" → "00:05"

            # Verificar quantos valores HH:MM existem em cel4_raw
            time_parts4 = [p for p in cel4_raw.split() if is_time(p)] if cel4_raw else []

            if len(time_parts4) >= 2:
                # ent2 e sai2 mesclados em [4]
                dia["ent2"] = time_parts4[0]
                dia["sai2"] = time_parts4[1]
                merged = True
            else:
                dia["ent2"] = cel4_raw if is_time(cel4_raw) else ""
                dia["sai2"] = cel5_raw if is_time(cel5_raw) else ""
                merged = False

            # ── Localizar CH e duração por valor semântico ────────────────────
            # CH é sempre um código de 5 dígitos (ex: "00021")
            # Duração é HH:MM imediatamente antes do CH
            # Percorre posições de maior para menor índice para pegar o CH mais à direita
            ch_positions = [9, 8, 7, 6, 5] if not merged else [7, 6, 5]

            found = False
            for ch_i in ch_positions:
                v = cel(ch_i)
                if is_ch(v):
                    dia["ch"] = v
                    prev = cel(ch_i - 1)
                    if is_time(prev):
                        dia["duracao"] = prev
                    # Formato B completo (CH em [9]): extrair ent3/sai3
                    if ch_i == 9:
                        dia["ent3"] = cel(6)
                        dia["sai3"] = cel(7)
                    found = True
                    break

            if not found:
                # CH ausente — determinar contexto pelo índice [9]
                # Formato B tem CH em [9] → duração em [8], ent3/sai3 em [6]/[7]
                # Formato A tem I/P/D em [9] → duração em [6], nada mais
                is_format_b = is_ch(cel(9))
                search_order = [8, 7, 6] if is_format_b else [6, 7, 8]

                for dur_i in search_order:
                    if is_time(cel(dur_i)):
                        dia["duracao"] = cel(dur_i)
                        if dur_i == 8 and is_format_b:
                            dia["ent3"] = cel(6)
                            dia["sai3"] = cel(7)
                        break

            dias.append(dia)

        return dias, horarios_contratuais

    # ──────────────────────────────────────────────
    # Extração MHTML
    # ──────────────────────────────────────────────

    def extrair_html_de_mhtml(self, caminho):
        """Abre um arquivo .mhtml e retorna o conteúdo HTML interno como string."""
        with open(caminho, "rb") as f:
            msg = email_lib.message_from_bytes(f.read())
        for part in msg.walk():
            if "html" in part.get_content_type():
                payload = part.get_payload(decode=True)
                charset = part.get_content_charset() or "utf-8"
                return payload.decode(charset, errors="replace")
        return None

    def extrair_info_funcionario_mhtml(self, soup):
        """Extrai dados do funcionário a partir do HTML parseado.
        Retorna (dict, motivo_falha). Células têm formato 'CAMPO:Valor'."""
        info = {}

        campo_map_lower = {
            "empresa"        : "empresa",
            "cnpj"           : "cnpj",
            "cei"            : "cei",
            "endereço"       : "endereco",
            "nome"           : "nome",
            "pis/pasep"      : "pis",
            "admissão"       : "admissao",
            "centro de custo": "centro_custo",
            "cpf"            : "cpf",
            "matrícula"      : "matricula",
            "departamento"   : "departamento",
            "cargo"          : "cargo",
        }

        for td in soup.find_all(["td", "th"]):
            texto = td.get_text(separator=" ", strip=True)
            if ":" not in texto:
                continue
            chave, _, valor = texto.partition(":")
            chave_norm = chave.strip().lower()
            valor = valor.strip()
            if chave_norm in campo_map_lower and valor:
                info[campo_map_lower[chave_norm]] = valor

        if not info.get("nome"):
            tem_empresa = "empresa" in info
            if not tem_empresa:
                motivo = "nenhum campo reconhecido — estrutura do MHTML diferente do esperado"
            else:
                motivo = "empresa identificada mas campo NOME não encontrado"
            return info, motivo

        return info, None

    def extrair_dados_ponto_mhtml(self, soup):
        """Extrai registros de ponto da tabela HTML.
        Cada linha-mestre tem 17 células: [0]=data, [3]=ENT1, [4]=SAÍ1,
        [5]=ENT2, [6]=SAÍ2. Linhas de detalhe (3 células) são ignoradas."""
        dias = []
        tables = soup.find_all("table")

        tabela_ponto = None
        for table in tables:
            primeira = table.find("tr")
            if primeira and "DIA" in primeira.get_text():
                tabela_ponto = table
                break

        if not tabela_ponto:
            return dias

        data_re = re.compile(r"(\d{2}/\d{2}/\d{2,4})")

        for tr in tabela_ponto.find_all("tr"):
            cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]

            if len(cells) < 17:
                continue
            m = data_re.match(cells[0])
            if not m:
                continue

            data_raw  = cells[0]
            data_part = m.group(1)

            try:
                from datetime import datetime as dt
                fmt = "%d/%m/%y" if len(data_part.split("/")[2]) == 2 else "%d/%m/%Y"
                data_fmt = dt.strptime(data_part, fmt).strftime("%d/%m/%Y")
            except ValueError:
                data_fmt = data_part

            dia_semana = data_raw.split("-")[-1].strip() if "-" in data_raw else ""

            ent1    = cells[3].strip()
            sai1    = cells[4].strip()
            ent2    = cells[5].strip()
            sai2    = cells[6].strip()
            duracao = cells[10].strip() if len(cells) > 10 else ""
            ch      = cells[12].strip() if len(cells) > 12 else ""

            # "marcacoes" = resumo visual das batidas (ex: "07:23 14:27")
            marcacoes = f"{ent1} {sai1}".strip() if ent1 or sai1 else ""

            dias.append({
                "dia"     : f"{data_fmt} - {dia_semana}",
                "marcacoes": marcacoes,
                "ent1"    : ent1,
                "sai1"    : sai1,
                "ent2"    : ent2,
                "sai2"    : sai2,
                "ent3"    : "",
                "sai3"    : "",
                "duracao" : duracao,
                "ch"      : ch,
            })

        return dias

    # ──────────────────────────────────────────────

    def extrair_funcionarios_de_pdf(self, pdf):
        """Processa um PDF com um ou mais funcionários, página a página.
        Retorna lista de tuplas (info_funcionario, dias)."""
        resultados = []

        texto_atual  = ""
        tabelas_atual = []

        def finalizar_secao(texto, tabelas):
            """Extrai e retorna (info, dias) de uma seção acumulada."""
            if not texto.strip():
                return None
            info, motivo = self.extrair_info_funcionario(texto)
            if not info.get("nome"):
                self.log(f"    ⚠️ Seção sem funcionário identificado: {motivo}", "error")
                return None
            dias, horarios = self.extrair_dados_ponto(tabelas)
            if horarios:
                info["horarios_contratuais"] = horarios
            return info, dias

        for pagina in pdf.pages:
            texto_pagina   = pagina.extract_text() or ""
            tabelas_pagina = []
            for t in (pagina.extract_tables() or []):
                if t:
                    tabelas_pagina.extend(t)

            # Nova seção começa quando a página contém um cabeçalho de funcionário
            e_novo_funcionario = bool(
                re.search(r'NOME:\s*.+?\s+PIS/PASEP:', texto_pagina, re.IGNORECASE)
            )

            if e_novo_funcionario and texto_atual:
                # Finaliza o funcionário anterior antes de começar o próximo
                resultado = finalizar_secao(texto_atual, tabelas_atual)
                if resultado:
                    resultados.append(resultado)
                texto_atual   = ""
                tabelas_atual = []

            texto_atual   += texto_pagina
            tabelas_atual += tabelas_pagina

        # Finaliza o último (ou único) funcionário
        if texto_atual:
            resultado = finalizar_secao(texto_atual, tabelas_atual)
            if resultado:
                resultados.append(resultado)

        return resultados

    def processar_pdfs_thread(self):
        """Thread principal de processamento."""
        self.processamento_ativo = True
        total_pdfs = len(self.pdf_paths)
        total_dias = 0
        todos_dias = []
        todos_funcionarios = {}  # chave: CPF → info_funcionario completo

        try:
            self.log(f"Iniciando processamento de {total_pdfs} arquivo(s)...", "processing")

            for idx, pdf_path in enumerate(self.pdf_paths):
                nome_arquivo = Path(pdf_path).name
                self.log(f"Processando: {nome_arquivo}", "processing")

                try:
                    if not Path(pdf_path).exists():
                        self.log(f"  ✗ Arquivo não encontrado: {nome_arquivo}", "error")
                        continue

                    ext = Path(pdf_path).suffix.lower()

                    if ext == ".mhtml":
                        html_str = self.extrair_html_de_mhtml(pdf_path)
                        if not html_str:
                            self.log(f"  ✗ Não foi possível extrair HTML de '{nome_arquivo}'.", "error")
                            continue

                        soup = BeautifulSoup(html_str, "html.parser")
                        info_funcionario, motivo_falha = self.extrair_info_funcionario_mhtml(soup)

                        if not info_funcionario.get("nome"):
                            self.log(f"  ⚠️ Nenhum funcionário identificado em '{nome_arquivo}'. "
                                     f"Motivo: {motivo_falha}", "error")
                            continue

                        dias_do_pdf = self.extrair_dados_ponto_mhtml(soup)

                    else:
                        try:
                            pdf_aberto = pdfplumber.open(pdf_path)
                        except Exception:
                            self.log(f"  ✗ Não foi possível abrir '{nome_arquivo}'. "
                                     f"O arquivo pode estar protegido por senha ou corrompido.", "error")
                            continue

                        with pdf_aberto as pdf:
                            funcionarios_no_pdf = self.extrair_funcionarios_de_pdf(pdf)

                        if not funcionarios_no_pdf:
                            self.log(f"  ✗ '{nome_arquivo}' não contém texto extraível ou "
                                     f"nenhum funcionário foi identificado.", "error")
                            continue

                        for info_funcionario, dias_do_pdf in funcionarios_no_pdf:
                            cpf = info_funcionario.get("cpf", info_funcionario.get("nome", ""))
                            for dia in dias_do_pdf:
                                dia["nome"]    = info_funcionario.get("nome", "")
                                dia["empresa"] = info_funcionario.get("empresa", "")
                                dia["cpf"]     = info_funcionario.get("cpf", "")

                            if cpf not in todos_funcionarios:
                                todos_funcionarios[cpf] = info_funcionario
                            todos_funcionarios[cpf].setdefault("_dias", []).extend(dias_do_pdf)

                            todos_dias.extend(dias_do_pdf)
                            total_dias += len(dias_do_pdf)
                            self.log(f"  ✓ {info_funcionario.get('nome', cpf)}: "
                                     f"{len(dias_do_pdf)} dia(s) extraído(s)", "success")
                        continue  # pula o bloco de associação abaixo (já feito acima)

                    # Associar dados básicos a cada dia e registrar funcionário completo (MHTML)
                    cpf = info_funcionario.get("cpf", info_funcionario.get("nome", ""))
                    for dia in dias_do_pdf:
                        dia["nome"]    = info_funcionario.get("nome", "")
                        dia["empresa"] = info_funcionario.get("empresa", "")
                        dia["cpf"]     = info_funcionario.get("cpf", "")

                    # Acumula dias do mesmo funcionário somando os que já existem
                    if cpf not in todos_funcionarios:
                        todos_funcionarios[cpf] = info_funcionario
                    todos_funcionarios[cpf].setdefault("_dias", []).extend(dias_do_pdf)

                    todos_dias.extend(dias_do_pdf)
                    total_dias += len(dias_do_pdf)
                    self.log(f"  ✓ {len(dias_do_pdf)} dia(s) extraído(s)", "success")

                except PermissionError:
                    self.log(f"  ✗ Sem permissão para ler '{nome_arquivo}'. "
                             f"Verifique se o arquivo está aberto em outro programa.", "error")
                except MemoryError:
                    self.log(f"  ✗ Memória insuficiente ao processar '{nome_arquivo}'. "
                             f"Tente processar menos arquivos por vez.", "error")
                except Exception as e:
                    self.log(f"  ✗ Erro inesperado em '{nome_arquivo}': {str(e)}", "error")

                self.root.after(0, lambda: self.progress_bar.set(((idx + 1) / total_pdfs)))

        except Exception as e:
            self.log(f"Erro crítico no processamento: {str(e)}", "error")
        finally:
            self.root.after(0, lambda: self._finalizar_processamento(
                todos_dias, total_pdfs, total_dias, todos_funcionarios))

    def salvar_excel(self, dados, todos_funcionarios):
        """Salva os dados em um arquivo Excel com cabeçalho, aba de resumo e nome inteligente."""

        # ── Nome do arquivo inteligente ─────────────────────────────────────
        nome_sugerido = self._gerar_nome_arquivo(dados, todos_funcionarios)

        excel_path = filedialog.asksaveasfilename(
            title="Salvar planilha Excel",
            defaultextension=".xlsx",
            initialfile=nome_sugerido,
            initialdir=self.dir_salvar_excel,
            filetypes=[("Arquivos Excel", "*.xlsx")]
        )

        if not excel_path:
            self.log("Operação cancelada pelo usuário.", "info")
            return

        self.dir_salvar_excel = str(Path(excel_path).parent)
        self.salvar_configuracoes(silencioso=True)

        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()
            ws = wb.active
            ws.title = "Ponto"

            funcionario_unico = len(todos_funcionarios) == 1

            # ── Estilos ─────────────────────────────────────────────────────
            fonte_cab    = Font(bold=True, color="FFFFFF", size=10)
            fill_cab     = PatternFill("solid", fgColor="1F4E79")
            fill_info    = PatternFill("solid", fgColor="D6E4F0")
            fonte_info   = Font(bold=True, size=10, color="1F4E79")
            fonte_label  = Font(bold=True, size=9,  color="2C3E50")
            alinhamento        = Alignment(vertical="center", horizontal="left")
            alinhamento_centro = Alignment(vertical="center", horizontal="center")
            borda_fina   = Border(
                bottom=Side(style="thin", color="BFBFBF"),
                top=Side(style="thin",    color="BFBFBF"),
            )

            linha_atual = 1

            # ── Helper: dias efetivamente trabalhados ────────────────────────
            def dias_trabalhados(dias_lista):
                return [d for d in dias_lista if any(
                    d.get(k, "").strip()
                    for k in ("ent1", "sai1", "ent2", "sai2", "ent3", "sai3")
                )]

            # ── Ordenação: nome → data (ANTES de escrever qualquer linha) ────
            def chave_ordenacao(dia):
                nome     = dia.get("nome", "")
                data_str = dia.get("dia", "").split(" ")[0]  # "DD/MM/AA" ou "DD/MM/AAAA"
                try:
                    from datetime import datetime as dt
                    partes = data_str.split("/")
                    fmt = "%d/%m/%y" if len(partes) == 3 and len(partes[2]) == 2 else "%d/%m/%Y"
                    data = dt.strptime(data_str, fmt)
                except ValueError:
                    data = datetime.min
                return (nome, data)

            dados = sorted(dados, key=chave_ordenacao)

            # ── Cabeçalho do funcionário (único) ─────────────────────────────
            if funcionario_unico:
                info = list(todos_funcionarios.values())[0]
                campos = [
                    ("Empresa",          info.get("empresa",      "")),
                    ("Nome",             info.get("nome",         "")),
                    ("CPF",              info.get("cpf",          "")),
                    ("PIS/PASEP",        info.get("pis",          "")),
                    ("Matrícula",        info.get("matricula",    "")),
                    ("Admissão",         info.get("admissao",     "")),
                    ("Cargo",            info.get("cargo",        "")),
                    ("Departamento",     info.get("departamento", "")),
                    ("Centro de Custo",  info.get("centro_custo", "")),
                ]
                campos = [(k, v) for k, v in campos if v]

                n_colunas_cab = 10  # DIA MARCAÇÕES ENT1 SAÍ1 ENT2 SAÍ2 ENT3 SAÍ3 DURAÇÃO CH

                ws.merge_cells(start_row=linha_atual, start_column=1,
                               end_row=linha_atual, end_column=n_colunas_cab)
                cel = ws.cell(row=linha_atual, column=1, value="DADOS DO FUNCIONÁRIO")
                cel.font      = fonte_cab
                cel.fill      = fill_cab
                cel.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[linha_atual].height = 20
                linha_atual += 1

                for i in range(0, len(campos), 2):
                    par = campos[i:i+2]
                    for col_offset, (label, valor) in enumerate(par):
                        col_l   = 1 + col_offset * 5
                        col_v   = col_l + 1
                        col_fim = min(col_l + 4, n_colunas_cab)

                        cel_l = ws.cell(row=linha_atual, column=col_l, value=label.upper() + ":")
                        cel_l.font      = fonte_label
                        cel_l.fill      = fill_info
                        cel_l.alignment = alinhamento

                        ws.merge_cells(start_row=linha_atual, start_column=col_v,
                                       end_row=linha_atual, end_column=col_fim)
                        cel_v = ws.cell(row=linha_atual, column=col_v, value=valor)
                        cel_v.font      = fonte_info
                        cel_v.fill      = fill_info
                        cel_v.alignment = alinhamento

                    ultimo_col_usado = 1 + len(par) * 5 - 1
                    for col_r in range(ultimo_col_usado + 1, n_colunas_cab + 1):
                        ws.cell(row=linha_atual, column=col_r).fill = fill_info

                    ws.row_dimensions[linha_atual].height = 18
                    linha_atual += 1

                # Separador
                for col_r in range(1, n_colunas_cab + 1):
                    ws.cell(row=linha_atual, column=col_r).fill = PatternFill("solid", fgColor="FFFFFF")
                ws.row_dimensions[linha_atual].height = 6
                linha_atual += 1

            # ── Colunas da tabela ─────────────────────────────────────────────
            if funcionario_unico:
                cabecalhos_visiveis = [
                    "DIA", "MARCAÇÕES", "ENT. 1", "SAÍ. 1",
                    "ENT. 2", "SAÍ. 2", "ENT. 3", "SAÍ. 3", "DURAÇÃO", "CH"
                ]
                valores_keys = [
                    "dia", "marcacoes", "ent1", "sai1",
                    "ent2", "sai2", "ent3", "sai3", "duracao", "ch"
                ]
                larguras = [18, 26, 10, 10, 10, 10, 10, 10, 10, 8]
            else:
                cabecalhos_visiveis = [
                    "FUNCIONÁRIO", "CPF", "DIA", "MARCAÇÕES", "ENT. 1", "SAÍ. 1",
                    "ENT. 2", "SAÍ. 2", "ENT. 3", "SAÍ. 3", "DURAÇÃO", "CH"
                ]
                valores_keys = [
                    "nome", "cpf", "dia", "marcacoes", "ent1", "sai1",
                    "ent2", "sai2", "ent3", "sai3", "duracao", "ch"
                ]
                larguras = [30, 16, 18, 26, 10, 10, 10, 10, 10, 10, 10, 8]

            # ── Fills alternados (usados no resumo e na tabela) ──────────────
            fill_par   = PatternFill("solid", fgColor="EBF5FB")
            fill_impar = PatternFill("solid", fgColor="FFFFFF")

            # ── Resumo ────────────────────────────────────────────────────────
            cab_resumo = [
                "NOME", "CPF", "PIS/PASEP", "CARGO", "ADMISSÃO",
                "MATRÍCULA", "DEPARTAMENTO", "CENTRO DE CUSTO",
                "DIAS TRABALHADOS", "MÉDIA HORAS/DIA"
            ]
            larg_resumo = [35, 16, 14, 25, 12, 12, 22, 18, 16, 16]

            def linha_resumo(info, dias_lista):
                trab        = dias_trabalhados(dias_lista)
                total_dias  = len(trab)
                total_horas = self._somar_duracoes([d.get("duracao", "") for d in trab])
                media       = self._media_horas(total_horas, total_dias)
                return [
                    info.get("nome",         ""),
                    info.get("cpf",          ""),
                    info.get("pis",          ""),
                    info.get("cargo",        ""),
                    info.get("admissao",     ""),
                    info.get("matricula",    ""),
                    info.get("departamento", ""),
                    info.get("centro_custo", ""),
                    total_dias,
                    media,
                ]

            fill_resumo_cab = PatternFill("solid", fgColor="2E86C1")

            if funcionario_unico:
                # Resumo inline entre cabeçalho e tabela
                ws.merge_cells(start_row=linha_atual, start_column=1,
                               end_row=linha_atual, end_column=n_colunas_cab)
                cel = ws.cell(row=linha_atual, column=1, value="RESUMO")
                cel.font      = fonte_cab
                cel.fill      = fill_resumo_cab
                cel.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[linha_atual].height = 20
                linha_atual += 1

                for col, cab in enumerate(cab_resumo, 1):
                    cel = ws.cell(row=linha_atual, column=col, value=cab)
                    cel.font      = fonte_cab
                    cel.fill      = fill_resumo_cab
                    cel.alignment = Alignment(horizontal="center", vertical="center")
                    cel.border    = borda_fina
                ws.row_dimensions[linha_atual].height = 18
                linha_atual += 1

                info_u  = list(todos_funcionarios.values())[0]
                vals_r  = linha_resumo(info_u, info_u.get("_dias", dados))
                for col, valor in enumerate(vals_r, 1):
                    cel = ws.cell(row=linha_atual, column=col, value=valor)
                    cel.fill      = fill_par
                    cel.alignment = alinhamento
                    cel.border    = borda_fina
                linha_atual += 1

                # Separador antes da tabela
                for col_r in range(1, n_colunas_cab + 1):
                    ws.cell(row=linha_atual, column=col_r).fill = PatternFill("solid", fgColor="FFFFFF")
                ws.row_dimensions[linha_atual].height = 6
                linha_atual += 1

            else:
                # Aba separada de Resumo para múltiplos funcionários
                ws_resumo = wb.create_sheet("Resumo")

                for col, cab in enumerate(cab_resumo, 1):
                    cel = ws_resumo.cell(row=1, column=col, value=cab)
                    cel.font      = fonte_cab
                    cel.fill      = fill_cab
                    cel.alignment = Alignment(horizontal="center", vertical="center")
                ws_resumo.row_dimensions[1].height = 20

                for row_r, (cpf, info) in enumerate(todos_funcionarios.items(), 2):
                    vals_r = linha_resumo(info, info.get("_dias", []))
                    fill   = fill_par if row_r % 2 == 0 else fill_impar
                    for col, valor in enumerate(vals_r, 1):
                        cel = ws_resumo.cell(row=row_r, column=col, value=valor)
                        cel.fill      = fill
                        cel.alignment = alinhamento
                        cel.border    = borda_fina

                for col, larg in enumerate(larg_resumo, 1):
                    ws_resumo.column_dimensions[get_column_letter(col)].width = larg

            # ── Cabeçalhos da tabela de ponto ────────────────────────────────
            for col, cab in enumerate(cabecalhos_visiveis, 1):
                cel = ws.cell(row=linha_atual, column=col, value=cab)
                cel.font      = fonte_cab
                cel.fill      = fill_cab
                cel.alignment = Alignment(horizontal="center", vertical="center")
                cel.border    = borda_fina
            ws.row_dimensions[linha_atual].height = 20
            linha_atual += 1

            # ── Dados de ponto ───────────────────────────────────────────────
            for i, dia in enumerate(dados):
                fill = fill_par if i % 2 == 0 else fill_impar
                col_dia = 1 if funcionario_unico else 3  # coluna DIA varia conforme o modo
                for col, key in enumerate(valores_keys, 1):
                    cel = ws.cell(row=linha_atual, column=col, value=dia.get(key, ""))
                    cel.fill      = fill
                    cel.alignment = alinhamento if col == col_dia else alinhamento_centro
                    cel.border    = borda_fina
                linha_atual += 1

            # ── Larguras das colunas ─────────────────────────────────────────
            for col, largura in enumerate(larguras, 1):
                ws.column_dimensions[get_column_letter(col)].width = largura

            wb.save(excel_path)
            self.log(f"Planilha salva: {excel_path}", "success")

            messagebox.showinfo(
                "Sucesso",
                f"Planilha Excel criada com sucesso!\n\n"
                f"Arquivo: {Path(excel_path).name}\n"
                f"Local: {Path(excel_path).parent}\n"
                f"Funcionários: {len(todos_funcionarios)}\n"
                f"Dias processados: {len(dados)}"
            )

        except Exception as e:
            self.log(f"Erro ao salvar Excel: {str(e)}", "error")
            messagebox.showerror("Erro", f"Não foi possível salvar a planilha:\n{str(e)}")

    def _gerar_nome_arquivo(self, dados, todos_funcionarios):
        """Gera o nome sugerido para o arquivo Excel com base no cenário."""
        MESES = {
            1:"Jan", 2:"Fev", 3:"Mar", 4:"Abr", 5:"Mai", 6:"Jun",
            7:"Jul", 8:"Ago", 9:"Set", 10:"Out", 11:"Nov", 12:"Dez"
        }

        def nome_curto(nome_completo):
            partes = nome_completo.strip().split()
            return f"{partes[0]}_{partes[-1]}" if len(partes) > 1 else partes[0]

        def mes_ano_dos_dias(dias_lista):
            """Detecta mês/ano predominante nas datas dos dias."""
            try:
                datas = [d.get("dia", "").split(" ")[0] for d in dias_lista if d.get("dia")]
                from collections import Counter
                meses_anos = []
                for data in datas:
                    partes = data.split("/")
                    if len(partes) == 3:
                        meses_anos.append((int(partes[1]), int(partes[2])))
                if not meses_anos:
                    return None
                mes, ano = Counter(meses_anos).most_common(1)[0][0]
                return f"{MESES[mes]}{ano}"
            except Exception:
                return None

        num_funcionarios = len(todos_funcionarios)

        if num_funcionarios == 1:
            info   = list(todos_funcionarios.values())[0]
            nc     = nome_curto(info.get("nome", "Funcionario"))
            dias_f = info.get("_dias", dados)

            if len(self.pdf_paths) == 1:
                # 1 arquivo, 1 funcionário → inclui mês
                mes_ano = mes_ano_dos_dias(dias_f)
                sufixo  = f"_{mes_ano}" if mes_ano else ""
                return f"Ponto_{nc}{sufixo}.xlsx"
            else:
                # N arquivos, mesmo funcionário → sem mês (ambíguo)
                return f"Ponto_{nc}.xlsx"
        else:
            # N funcionários diferentes → contagem + timestamp do mês atual
            mes_ano = MESES[datetime.now().month] + str(datetime.now().year)
            return f"Ponto_{num_funcionarios}Funcionarios_{mes_ano}.xlsx"

    def _somar_duracoes(self, duracoes):
        """Soma uma lista de strings 'HH:MM' e retorna o total em 'HH:MM'."""
        total_min = 0
        for d in duracoes:
            if not d or ":" not in d:
                continue
            try:
                h, m = d.strip().split(":")
                total_min += int(h) * 60 + int(m)
            except ValueError:
                continue
        return f"{total_min // 60:02d}:{total_min % 60:02d}"

    def _media_horas(self, total_horas, total_dias):
        """Calcula a média de horas por dia a partir de 'HH:MM' e quantidade de dias."""
        if not total_dias or ":" not in total_horas:
            return "00:00"
        try:
            h, m = total_horas.split(":")
            total_min = int(h) * 60 + int(m)
            media_min = total_min // total_dias
            return f"{media_min // 60:02d}:{media_min % 60:02d}"
        except ValueError:
            return "00:00"

    def separar_pdf_por_funcionario(self):
        """Abre um PDF com múltiplos funcionários e gera um ZIP com um PDF por funcionário."""
        pdf_path = filedialog.askopenfilename(
            title="Selecionar PDF com múltiplos funcionários",
            filetypes=[("Arquivos PDF", "*.pdf")],
            initialdir=self.dir_abrir_pdf
        )
        if not pdf_path:
            return

        self.dir_abrir_pdf = str(Path(pdf_path).parent)
        self.salvar_configuracoes(silencioso=True)

        zip_path = filedialog.asksaveasfilename(
            title="Salvar ZIP com PDFs separados",
            defaultextension=".zip",
            initialfile=f"{Path(pdf_path).stem}_separados.zip",
            initialdir=self.dir_salvar_excel,
            filetypes=[("Arquivo ZIP", "*.zip")]
        )
        if not zip_path:
            return

        self.log("Iniciando separação de PDF por funcionário...", "processing")
        import threading
        t = threading.Thread(
            target=self._separar_pdf_worker,
            args=(pdf_path, zip_path),
            daemon=True
        )
        t.start()

    def _separar_pdf_worker(self, pdf_path, zip_path):
        """Thread que separa o PDF e gera o ZIP."""
        try:
            # Detectar grupos de páginas por funcionário
            grupos = []
            grupo_atual = None

            with pdfplumber.open(pdf_path) as pdf:
                for i, pagina in enumerate(pdf.pages):
                    texto = pagina.extract_text() or ""
                    match = re.search(r'NOME:\s*(.+?)\s+PIS/PASEP:', texto, re.IGNORECASE)
                    if match:
                        nome = match.group(1).strip()
                        grupo_atual = {"nome": nome, "paginas": [i]}
                        grupos.append(grupo_atual)
                    elif grupo_atual:
                        grupo_atual["paginas"].append(i)

            if not grupos:
                self.root.after(0, lambda: (
                    self.log("⚠️ Nenhum funcionário identificado no PDF.", "error"),
                    messagebox.showwarning("Atenção", "Nenhum funcionário identificado no PDF.")
                ))
                return

            if len(grupos) == 1:
                self.root.after(0, lambda: (
                    self.log("⚠️ O PDF contém apenas um funcionário — separação desnecessária.", "error"),
                    messagebox.showwarning("Atenção", "O PDF contém apenas um funcionário.")
                ))
                return

            self.log(f"  {len(grupos)} funcionário(s) detectado(s). Gerando PDFs...", "processing")

            # Gerar um PDF por funcionário e compactar no ZIP
            reader = pypdf.PdfReader(pdf_path)

            def nome_para_arquivo(nome):
                """Sanitiza o nome para uso em nome de arquivo."""
                import unicodedata
                nfkd = unicodedata.normalize("NFKD", nome)
                ascii_nome = nfkd.encode("ASCII", "ignore").decode()
                return re.sub(r"[^\w\s-]", "", ascii_nome).strip().replace(" ", "_")

            with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
                nomes_usados = {}
                for grupo in grupos:
                    writer = pypdf.PdfWriter()
                    for idx_pagina in grupo["paginas"]:
                        writer.add_page(reader.pages[idx_pagina])

                    nome_base = nome_para_arquivo(grupo["nome"])

                    # Evitar colisão de nomes
                    if nome_base in nomes_usados:
                        nomes_usados[nome_base] += 1
                        nome_arquivo = f"{nome_base}_{nomes_usados[nome_base]}.pdf"
                    else:
                        nomes_usados[nome_base] = 1
                        nome_arquivo = f"{nome_base}.pdf"

                    import io
                    buffer = io.BytesIO()
                    writer.write(buffer)
                    zf.writestr(nome_arquivo, buffer.getvalue())

                    self.log(f"  ✓ {grupo['nome']} → {nome_arquivo} "
                             f"({len(grupo['paginas'])} página(s))", "success")

            self.root.after(0, lambda: (
                self.log(f"\n✅ ZIP gerado com sucesso: {zip_path}", "success"),
                messagebox.showinfo(
                    "Concluído",
                    f"PDFs separados com sucesso!\n\n"
                    f"Funcionários: {len(grupos)}\n"
                    f"Arquivo: {Path(zip_path).name}\n"
                    f"Local: {Path(zip_path).parent}"
                )
            ))

        except PermissionError:
            self.root.after(0, lambda: (
                self.log("✗ Sem permissão para ler o PDF ou salvar o ZIP.", "error"),
                messagebox.showerror("Erro", "Sem permissão para acessar o arquivo.")
            ))
        except Exception as e:
            self.root.after(0, lambda: (
                self.log(f"✗ Erro ao separar PDF: {str(e)}", "error"),
                messagebox.showerror("Erro", f"Erro ao separar PDF:\n{str(e)}")
            ))

    def iniciar_processamento(self):
        """Inicia o processamento em thread separada."""
        if not self.pdf_paths:
            messagebox.showwarning("Atenção", "Selecione pelo menos um arquivo.")
            return

        if self.processamento_ativo:
            return

        # Confirmar processamento
        resposta = messagebox.askyesno(
            "Confirmar Processamento",
            f"Deseja processar {len(self.pdf_paths)} arquivo(s) PDF?\n\n"
            "O processo pode levar alguns segundos dependendo do tamanho dos arquivos."
        )

        if not resposta:
            return

        self.btn_processar.configure(state="disabled")
        self.log("--- Iniciando processamento ---", "processing")
        self.progress_var.set(0)

        # Executar em thread
        thread = threading.Thread(target=self.processar_pdfs_thread, daemon=True)
        thread.start()

    def _config_path(self):
        """Retorna o caminho do config.json em %APPDATA%\\ControlIDReader\\."""
        appdata = Path(os.environ.get("APPDATA", Path.home() / "AppData" / "Roaming"))
        config_dir = appdata / "ControlIDReader"
        config_dir.mkdir(parents=True, exist_ok=True)
        return config_dir / "config.json"

    def carregar_configuracoes(self):
        """Carrega configurações salvas."""
        try:
            config_file = self._config_path()
            if config_file.exists():
                with open(config_file, 'r', encoding='utf-8') as f:
                    content = f.read().strip()
                    if content:
                        config = json.loads(content)
                        if config.get("dir_abrir_pdf") and Path(config["dir_abrir_pdf"]).exists():
                            self.dir_abrir_pdf = config["dir_abrir_pdf"]
                        if config.get("dir_salvar_excel") and Path(config["dir_salvar_excel"]).exists():
                            self.dir_salvar_excel = config["dir_salvar_excel"]
                        if config.get("tema") in ("dark", "light"):
                            self._tema_atual = config["tema"]
                            ctk.set_appearance_mode(self._tema_atual)
                            label = "🌙  Dark" if self._tema_atual == "dark" else "☀️  Light"
                            self.btn_tema.configure(text=label)
                            self._atualizar_cores_log()
        except Exception:
            pass  # Ignora erros de configuração — defaults já estão definidos no __init__

    def salvar_configuracoes(self, silencioso=False):
        """Salva configurações atuais."""
        try:
            config = {
                "dir_abrir_pdf": self.dir_abrir_pdf,
                "dir_salvar_excel": self.dir_salvar_excel,
                "tema": self._tema_atual,
                "data_ultima_execucao": datetime.now().isoformat()
            }
            config_file = self._config_path()
            with open(config_file, 'w', encoding='utf-8') as f:
                json.dump(config, f, indent=2, ensure_ascii=False)
            if not silencioso:
                self.log("Configurações salvas com sucesso.", "success")
                messagebox.showinfo("Sucesso", "Configurações salvas!")
        except Exception as e:
            self.log(f"Erro ao salvar configurações: {e}", "error")
            if not silencioso:
                messagebox.showerror("Erro", f"Não foi possível salvar configurações:\n{str(e)}")


def main():
    """Função principal."""
    root = ctk.CTk()
    app = PdfToExcelApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
