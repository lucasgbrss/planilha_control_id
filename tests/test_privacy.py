import shutil
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path

from openpyxl import Workbook, load_workbook

from control_id_reader.privacy import (
    PrivacyError,
    anonymize_excel_file,
    restore_excel_file,
    suggest_anonymized_path,
    suggest_key_path,
    suggest_restored_path,
)


@contextmanager
def temp_workspace_dir():
    root = Path.cwd() / "test_output_privacy"
    root.mkdir(exist_ok=True)
    path = root / f"case_{uuid.uuid4().hex}"
    path.mkdir()
    try:
        yield str(path)
    finally:
        shutil.rmtree(path, ignore_errors=True)


class PrivacyTest(unittest.TestCase):
    def test_suggest_privacy_paths(self):
        path = Path("Ponto_Ana.xlsx")

        self.assertEqual(suggest_anonymized_path(path).name, "Ponto_Ana_anonimizado.xlsx")
        self.assertEqual(suggest_key_path(path).name, "Ponto_Ana.cidkey")
        self.assertEqual(suggest_restored_path(Path("Ponto_Ana_anonimizado.xlsx")).name, "Ponto_Ana_restaurado.xlsx")

    def test_anonymize_and_restore_excel_file(self):
        with temp_workspace_dir() as tmp:
            tmp_path = Path(tmp)
            original = tmp_path / "ponto.xlsx"
            anonimizado = tmp_path / "ponto_anonimizado.xlsx"
            restaurado = tmp_path / "ponto_restaurado.xlsx"
            chave = tmp_path / "ponto.cidkey"

            wb = Workbook()
            ws = wb.active
            ws.title = "Ponto"
            ws.append(["FUNCIONÁRIO", "CPF", "DIA", "DURAÇÃO", "CH", "OBS"])
            ws.append([
                "Ana Souza",
                "123.456.789-01",
                "05/04/26 - SEG",
                "08:00",
                "00021",
                "Ana Souza possui ajuste manual.",
            ])
            resumo = wb.create_sheet("Resumo")
            resumo.append(["NOME", "CPF", "PIS/PASEP", "MATRÍCULA", "DIAS TRABALHADOS"])
            resumo.append(["Ana Souza", "123.456.789-01", "987654321", "42", 1])
            wb.save(original)

            result = anonymize_excel_file(original, anonimizado, chave, "senha-forte")

            self.assertGreaterEqual(result["replacements"], 6)
            wb_anon = load_workbook(anonimizado)
            self.assertEqual(wb_anon["Ponto"]["A2"].value, "FUNCIONARIO_001")
            self.assertEqual(wb_anon["Ponto"]["B2"].value, "CPF_001")
            self.assertEqual(wb_anon["Ponto"]["C2"].value, "05/04/26 - SEG")
            self.assertEqual(wb_anon["Ponto"]["D2"].value, "08:00")
            self.assertEqual(wb_anon["Ponto"]["E2"].value, "00021")
            self.assertEqual(wb_anon["Ponto"]["F2"].value, "FUNCIONARIO_001 possui ajuste manual.")
            self.assertNotIn("Ana Souza", chave.read_text(encoding="utf-8"))

            restore_excel_file(anonimizado, restaurado, chave, "senha-forte")

            wb_rest = load_workbook(restaurado)
            self.assertEqual(wb_rest["Ponto"]["A2"].value, "Ana Souza")
            self.assertEqual(wb_rest["Ponto"]["B2"].value, "123.456.789-01")
            self.assertEqual(wb_rest["Ponto"]["F2"].value, "Ana Souza possui ajuste manual.")
            self.assertEqual(wb_rest["Resumo"]["D2"].value, "42")

    def test_restore_rejects_wrong_password(self):
        with temp_workspace_dir() as tmp:
            tmp_path = Path(tmp)
            original = tmp_path / "ponto.xlsx"
            anonimizado = tmp_path / "ponto_anonimizado.xlsx"
            restaurado = tmp_path / "ponto_restaurado.xlsx"
            chave = tmp_path / "ponto.cidkey"

            wb = Workbook()
            ws = wb.active
            ws.append(["FUNCIONÁRIO", "CPF"])
            ws.append(["Ana Souza", "123.456.789-01"])
            wb.save(original)

            anonymize_excel_file(original, anonimizado, chave, "senha-certa")

            with self.assertRaises(PrivacyError):
                restore_excel_file(anonimizado, restaurado, chave, "senha-errada")

    def test_summary_block_does_not_anonymize_punch_table_dates(self):
        with temp_workspace_dir() as tmp:
            tmp_path = Path(tmp)
            original = tmp_path / "ponto.xlsx"
            anonimizado = tmp_path / "ponto_anonimizado.xlsx"
            chave = tmp_path / "ponto.cidkey"

            wb = Workbook()
            ws = wb.active
            ws.title = "Ponto"
            ws.append(["NOME", "CPF", "DIAS TRABALHADOS"])
            ws.append(["Ana Souza", "123.456.789-01", 1])
            ws.append([])
            ws.append(["DIA", "MARCAÇÕES", "CH"])
            ws.append(["05/04/26 - SEG", "08:00 17:00", "00021"])
            wb.save(original)

            anonymize_excel_file(original, anonimizado, chave, "senha-forte")

            wb_anon = load_workbook(anonimizado)
            self.assertEqual(wb_anon["Ponto"]["A2"].value, "FUNCIONARIO_001")
            self.assertEqual(wb_anon["Ponto"]["A5"].value, "05/04/26 - SEG")


if __name__ == "__main__":
    unittest.main()
